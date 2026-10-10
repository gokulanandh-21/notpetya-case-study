import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, col_widths, header_bg="1A2536", alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for j, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            if is_header:
                set_cell_background(cell, header_bg)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
                        run.font.name = "Calibri"
            else:
                bg = alt_bg if i % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.font.name = "Calibri"
                        run.font.color.rgb = RGBColor(30, 41, 59)

def build_report():
    doc = Document()

    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # TITLE & SUBTITLE
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("THE NOTPETYA CYBER WAR:\nRELATIONAL DATABASE-DRIVEN INCIDENT INVESTIGATION PLATFORM")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Database Management Systems (DBMS) Academic Case Study & Implementation Report")
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph() # Spacer

    # METADATA TABLE
    meta_table = doc.add_table(rows=5, cols=2)
    meta_data = [
        ("Course / Subject:", "Database Management Systems (DBMS) Laboratory & Project"),
        ("Project Title:", "NotPetya Cyber Incident: Relational 3NF Threat Intelligence & Query Engine"),
        ("Live Demonstration URL:", "http://notpetya.vercel.app/"),
        ("GitHub Repository:", "https://github.com/gokulanandh-21/notpetya-case-study"),
        ("Primary Database Engine:", "PostgreSQL 14+ (Hosted on Neon / Prisma Cloud, Managed via node-postgres)")
    ]
    for idx, (label, val) in enumerate(meta_data):
        meta_table.rows[idx].cells[0].paragraphs[0].add_run(label).bold = True
        meta_table.rows[idx].cells[1].paragraphs[0].add_run(val)
    style_table(meta_table, [2.2, 4.3], header_bg="334155")

    doc.add_page_break()

    # SECTION 1: EXECUTIVE SUMMARY
    h1 = doc.add_heading("1. Executive Summary & Project Objectives", level=1)
    h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "On June 27, 2017, the NotPetya cyber incident emerged as one of the most economically devastating cyberattacks in recorded history. "
        "Originating from a compromised software update mechanism of Ukraine's ubiquitous tax software (M.E.Doc), the malware weaponized NSA-developed exploits "
        "(EternalBlue, EternalRomance) combined with memory credential scraping (Mimikatz) and administrative execution utilities (PsExec, WMIC). "
        "Despite presenting as ransomware demanding $300 in Bitcoin, the payload was a military-grade destructive wiper that irreversibly wiped the Master Boot Record (MBR) "
        "and encrypted the Master File Table (MFT) using non-recoverable keys, inflicting over $10 Billion USD in verified global damages across multinational enterprises such as "
        "A.P. Møller–Maersk, Merck & Co., and FedEx TNT Express."
    )
    doc.add_paragraph(
        "The objective of this academic project is to bridge Database Management Systems (DBMS) theory with Cyber Threat Intelligence (CTI). "
        "Rather than presenting static text or simulated UI data, this system is backed by a fully normalized (Third Normal Form - 3NF) PostgreSQL relational database. "
        "The web platform enables students, faculty, and security analysts to interactively explore and query the NotPetya forensic corpus via two synchronized modalities: "
        "(1) a Natural Language Query Engine powered by automated SQL translation, and (2) a secure, interactive manual SQL Console enforcing strict read-only execution policies."
    )

    # SECTION 2: RELATIONAL DATABASE ARCHITECTURE & 3NF
    h2 = doc.add_heading("2. Relational Database Design & Normalization (3NF)", level=1)
    h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "To ensure zero data redundancy, high transactional integrity, and prevent update/deletion anomalies, the schema is organized in strict Third Normal Form (3NF):\n"
        "• First Normal Form (1NF): All attributes contain strictly atomic values; multi-valued attributes (such as aliases) are decoupled or handled as typed relational arrays.\n"
        "• Second Normal Form (2NF): All relations possess explicit primary keys, and every non-key attribute is fully functionally dependent on the entire primary key.\n"
        "• Third Normal Form (3NF): Non-prime attributes depend solely on candidate keys with zero transitive dependencies. Junction tables resolve all Many-to-Many (M:N) cardinalities."
    )

    doc.add_heading("2.1 Entity Relationship Table Directory", level=2)
    tbl_summary = doc.add_table(rows=16, cols=4)
    tbl_headers = ["Table Name", "Relation Role", "Primary Key (PK)", "Foreign Key (FK) References"]
    for j, h in enumerate(tbl_headers):
        tbl_summary.rows[0].cells[j].paragraphs[0].text = h

    table_records = [
        ("incidents", "Master Root Entity", "incident_id", "None (Top-level parent)"),
        ("sources", "Evidence / Citations", "source_id", "None (Independent registry)"),
        ("attack_stages", "Kill Chain Phases", "stage_id", "None (Ordered sequentially)"),
        ("attack_techniques", "MITRE ATT&CK Matrix", "technique_id", "source_id → sources"),
        ("timeline_events", "Chronological Log", "event_id", "incident_id, source_id"),
        ("malware_components", "Modular Payloads", "component_id", "incident_id, source_id"),
        ("vulnerabilities", "CVE Vulnerabilities", "vulnerability_id", "source_id → sources"),
        ("victims", "Afflicted Entities", "victim_id", "source_id → sources"),
        ("impact_events", "Forensic Consequences", "impact_id", "victim_id, source_id"),
        ("mitigations", "Defensive Controls", "mitigation_id", "source_id → sources"),
        ("incident_techniques", "Junction (M:N)", "(incident_id, technique_id)", "incidents, attack_techniques"),
        ("technique_vulnerabilities", "Junction (M:N)", "(technique_id, vulnerability_id)", "attack_techniques, vulnerabilities"),
        ("incident_victims", "Junction (M:N)", "(incident_id, victim_id)", "incidents, victims"),
        ("stage_techniques", "Junction (M:N)", "(stage_id, technique_id)", "attack_stages, attack_techniques"),
        ("mitigation_techniques", "Junction (M:N)", "(mitigation_id, technique_id)", "mitigations, attack_techniques")
    ]

    for i, row in enumerate(table_records):
        for j, val in enumerate(row):
            tbl_summary.rows[i+1].cells[j].paragraphs[0].text = val
    style_table(tbl_summary, [1.5, 1.4, 1.6, 2.0])

    doc.add_page_break()

    # SECTION 3: SCHEMA DDL & DATA INTEGRITY
    h3 = doc.add_heading("3. Schema DDL & Referential Integrity Constraints", level=1)
    h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Below is an excerpt of the core Data Definition Language (DDL) demonstrating referential constraints, "
        "CHECK validation constraints, and indexing structures implemented in PostgreSQL:"
    )

    code_p = doc.add_paragraph()
    code_run = code_p.add_run(
        "-- Sources Table with Credibility CHECK Constraints\n"
        "CREATE TABLE sources (\n"
        "    source_id VARCHAR(32) PRIMARY KEY,\n"
        "    source_name VARCHAR(255) NOT NULL,\n"
        "    publisher VARCHAR(255) NOT NULL,\n"
        "    source_type VARCHAR(64) NOT NULL CHECK (source_type IN (\n"
        "        'Government / Advisory', 'Threat Intelligence Report',\n"
        "        'Vendor Technical Bulletin', 'MITRE ATT&CK',\n"
        "        'Corporate Disclosure', 'Reputable Journalism')),\n"
        "    publication_date DATE,\n"
        "    url TEXT,\n"
        "    credibility_level VARCHAR(32) NOT NULL DEFAULT 'High' CHECK (\n"
        "        credibility_level IN ('Primary / Official', 'High', 'Corroborated Secondary'))\n"
        ");\n\n"
        "-- Attack Techniques Aligned with MITRE ATT&CK\n"
        "CREATE TABLE attack_techniques (\n"
        "    technique_id VARCHAR(32) PRIMARY KEY,\n"
        "    mitre_id VARCHAR(32) NOT NULL,\n"
        "    technique_name VARCHAR(255) NOT NULL,\n"
        "    tactic VARCHAR(128) NOT NULL,\n"
        "    description TEXT NOT NULL,\n"
        "    observed_behavior TEXT NOT NULL,\n"
        "    detection_opportunity TEXT,\n"
        "    mitigation_strategy TEXT,\n"
        "    source_id VARCHAR(32) REFERENCES sources(source_id) ON DELETE SET NULL\n"
        ");\n"
        "CREATE INDEX idx_attack_techniques_tactic ON attack_techniques(tactic);\n"
        "CREATE INDEX idx_attack_techniques_mitre ON attack_techniques(mitre_id);\n\n"
        "-- Impact Events Distinguishing Direct Actions from Consequences\n"
        "CREATE TABLE impact_events (\n"
        "    impact_id VARCHAR(32) PRIMARY KEY,\n"
        "    victim_id VARCHAR(32) NOT NULL REFERENCES victims(victim_id) ON DELETE CASCADE,\n"
        "    impact_type VARCHAR(128) NOT NULL,\n"
        "    severity VARCHAR(32) NOT NULL CHECK (severity IN ('Catastrophic', 'Critical', 'Severe', 'Moderate')),\n"
        "    direct_malware_action TEXT NOT NULL,\n"
        "    operational_consequence TEXT NOT NULL,\n"
        "    recovery_duration_days INT,\n"
        "    source_id VARCHAR(32) REFERENCES sources(source_id) ON DELETE SET NULL\n"
        ");"
    )
    code_run.font.name = "Consolas"
    code_run.font.size = Pt(8.5)
    code_run.font.color.rgb = RGBColor(15, 23, 42)

    # SECTION 4: THE DUAL QUERY ENGINE & SECURITY
    h4 = doc.add_heading("4. The Dual Query Engine & Security Architecture", level=1)
    h4.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "A critical innovation of this project is the Dual Query Engine architecture. Users can interact with the underlying relational database "
        "through two safe, pedagogical workflows without risking SQL injection or unauthorized schema modification."
    )

    doc.add_heading("4.1 Natural Language Query to SQL Flow", level=2)
    doc.add_paragraph(
        "1. Natural Language Input: User enters an investigative question (e.g., 'Show all lateral movement techniques').\n"
        "2. Automated Translation (/api/nl-query): The system prompts an LLM compiler constrained strictly by the database catalog, "
        "with an instant fallback to a deterministic cyber regex rules compiler if offline.\n"
        "3. AST & Lexical Policy Guard: The generated SQL is parsed by the security validator.\n"
        "4. Execution & Display: Generated SQL is populated into an editable syntax console and executed against PostgreSQL (/api/query)."
    )

    doc.add_heading("4.2 Read-Only Enforcement & Anti-SQL Injection Defenses", level=2)
    doc.add_paragraph(
        "Because the query console is exposed to user input, multiple layers of defense-in-depth are enforced:\n"
        "• Strict Keyword Filter: Queries must begin strictly with SELECT or WITH. Any DDL/DML token (INSERT, UPDATE, DELETE, DROP, ALTER, TRUNCATE, GRANT, REVOKE, PG_SLEEP, INTO) is rejected prior to execution.\n"
        "• Anti-Stacking Semicolon Guard: Trailing semicolons are sanitized, and internal semicolons are rejected to block stacked injection attacks.\n"
        "• Server-Side Execution Timeout: The PostgreSQL client executes `SET statement_timeout = 5000` on every pooled connection to prevent denial-of-service via computationally expensive queries.\n"
        "• Output Row Cap: Queries are automatically capped at 100 rows using automatic LIMIT injection.\n"
        "• Pedagogical Error Diagnostics: When typographical errors occur (such as `SELECT * FORM victims;`), the system outputs friendly instructional suggestions ('Did you mean FROM?') rather than exposing backend database stack traces."
    )

    doc.add_page_break()

    # SECTION 5: BENCHMARK SQL QUERIES & DEMONSTRATION
    h5 = doc.add_heading("5. Academic Benchmark SQL Queries", level=1)
    h5.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "The system includes 8 academic benchmark queries that verify core relational algebra operations (Selection, Projection, Cartesian Product, Joins, Grouping, and Aggregation):"
    )

    demo_queries = [
        ("Benchmark 1: Full Relation Selection", "SELECT * FROM attack_techniques;", "Verifies basic relational scanning across all 10 MITRE techniques."),
        ("Benchmark 2: Filtered Projection (WHERE Clause)", "SELECT technique_name, mitre_id FROM attack_techniques WHERE tactic = 'Lateral Movement';", "Filters propagation techniques (EternalBlue, PsExec, WMIC)."),
        ("Benchmark 3: Entity Inspection (Vulnerabilities)", "SELECT cve_id, ms_bulletin, vulnerability_name, cvss_score FROM vulnerabilities;", "Retrieves CVE-2017-0144 and CVE-2017-0145 details with CVSS scores."),
        ("Benchmark 4: Geographic & Sector Projections", "SELECT organization_name, country, sector FROM victims ORDER BY country;", "Projects victim demographic profiles across multinational sectors."),
        ("Benchmark 5: Aggregation & Grouping (GROUP BY)", "SELECT impact_type, COUNT(*) as count FROM impact_events GROUP BY impact_type;", "Demonstrates relational aggregation across distinct operational outage categories."),
        ("Benchmark 6: Sorting & Ranking (ORDER BY DESC)", "SELECT tactic, COUNT(*) as technique_count FROM attack_techniques GROUP BY tactic ORDER BY technique_count DESC;", "Ranks MITRE tactics by the volume of weaponized techniques."),
        ("Benchmark 7: 2-Table Relational Inner Join", "SELECT v.organization_name, i.impact_type, i.severity FROM victims v JOIN impact_events i ON v.victim_id = i.victim_id;", "Demonstrates primary-to-foreign key joining across victim and impact tables."),
        ("Benchmark 8: Multi-Table Join via M:N Junction Table", "SELECT a.technique_name, s.stage_name FROM attack_techniques a JOIN stage_techniques st ON a.technique_id = st.technique_id JOIN attack_stages s ON st.stage_id = s.stage_id ORDER BY s.stage_number;", "Traverses a 3-table relational path across the M:N junction table stage_techniques.")
    ]

    for title, sql, desc in demo_queries:
        doc.add_heading(title, level=2)
        doc.add_paragraph(desc)
        p_sql = doc.add_paragraph()
        r_sql = p_sql.add_run(sql)
        r_sql.font.name = "Consolas"
        r_sql.font.size = Pt(9.5)
        r_sql.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_page_break()

    # SECTION 6: FORENSIC RESEARCH AUDIT & CORRECTIONS
    h6 = doc.add_heading("6. Forensic Fact-Checking & Academic Audit", level=1)
    h6.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "In accordance with rigorous academic standards, all historical and technical claims were cross-referenced against official disclosures "
        "(CISA TA17-181A, UK NCSC, US Department of Justice 2020 Indictment, Microsoft MSRC, and corporate SEC Form 10-K filings). "
        "Key corrections applied to the investigation corpus include:"
    )

    doc.add_paragraph(
        "1. Removal of TradeLens Anachronism: Unsubstantiated claims attributing the disruption of TradeLens to the 2017 NotPetya attack were excised. "
        "Historically accurate descriptions ('terminal cargo booking and operational logistics systems') were instituted.\n"
        "2. Code-Signing Precision: Statements claiming that NotPetya itself was signed with a legitimate digital certificate were corrected. "
        "The update mechanism of M.E.Doc was compromised to distribute the malicious DLL; the wiper binary perfc.dat was unsigned.\n"
        "3. Removal of Invented Minute Timestamps: Unsupported precise minute-by-minute timestamps (e.g., 10:30, 10:45) were replaced with qualified operational phase bands ('Morning (~10:30 UTC)', 'Midday', 'Early Afternoon').\n"
        "4. Qualified Ukrainian Saturation Claims: Blanket claims stating '80% of Ukrainian businesses were infected' were adjusted to accurate legal-regulatory context: "
        "'M.E.Doc was mandatory tax software utilized by the vast majority of corporations operating in Ukraine'.\n"
        "5. Technical Disaggregation of Impact: Phrasing stating 'Domain Controllers destroyed' was refined to 'Active Directory services rendered completely unavailable across global sites', "
        "rigorously separating direct disk encryption actions from systemic operational enterprise consequences."
    )

    # SECTION 7: SOC SIMULATOR & DEFENSE-IN-DEPTH
    h7 = doc.add_heading("7. Interactive SOC Simulator & Kill-Chain Defense", level=1)
    h7.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "To complement the DBMS architecture, two advanced threat intelligence modules were integrated:\n"
        "• 'You Are the SOC' Simulator: A 3-phase real-time incident decision matrix challenging analysts with authentic telemetry "
        "(C2 HTTP beacons, EternalBlue SMB sweeps, and LSASS injection/PsExec service creation). Evaluates containment outcomes and teaches defensive consequences.\n"
        "• What Would Have Stopped NotPetya?: A 6-tier defense-in-depth breakdown covering (1) Software Supply Chain SLSA attestation, "
        "(2) Windows Defender Credential Guard, (3) SMBv1 removal, (4) Active Directory Tiered Administration, (5) Workstation microsegmentation, and (6) Air-gapped immutable disaster recovery."
    )

    # SECTION 8: CONCLUSION
    h8 = doc.add_heading("8. Conclusion & Demonstration Verification", level=1)
    h8.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "The upgraded NotPetya Investigation Platform demonstrates that complex cybersecurity case studies can be elevated from static web pages "
        "into academically rigorous, database-backed laboratory environments. By integrating a normalized 3NF PostgreSQL database, a dual natural language / SQL execution engine, "
        "and strict read-only execution guardrails, this project serves as a comprehensive demonstration of DBMS principles in high-stakes threat intelligence applications."
    )

    # OUTPUT PATH
    out_dir = r"C:\Users\Gokul\Documents\notpetya-deploy"
    out_path = os.path.join(out_dir, "NOTPETYA_DBMS_CASE_STUDY_PROJECT_REPORT.docx")
    doc.save(out_path)
    print(f"Report generated successfully at: {out_path}")

if __name__ == "__main__":
    build_report()
